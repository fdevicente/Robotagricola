# -*- coding: utf-8 -*-
"""Los botones salen del Master, no van escritos a mano.

Las labores cambian con la temporada: hoy poda, en marzo cosecha. Una lista fija
envejece sin que nadie lo note y Juan termina apretando "Otra..." siempre.
"""
from openpyxl import Workbook

from modules.opciones_capataz import labores_frecuentes, maquinas_para_botones


def _excel_maquinas(tmp_path):
    """Reproduce el caso real: tractores que se usan vs fichas cargadas una vez."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws.append(["Fecha", "Hora", "Tipo", "Actividad", "Cultivo", "Sector",
               "Jornadas Hombre", "Trabajadores", "Insumo", "Cantidad",
               "Unidad", "Registro", "Registrado por", "Máquina", "Odómetro",
               "Horas Día", "Superficie ha", "Días Cubiertos"])

    def lectura(fecha, maquina, odo):
        ws.append([fecha, "10:00", "MAQUINARIA", "Lectura de horómetro",
                   "GENERAL", "", None, "", "", None, "", "t", "Juan Parada",
                   maquina, odo, None, None, None])

    for i in range(10):                        # el que mas se usa
        lectura("2026-07-%02d" % (i + 1), "TRACTOR JOHN DEERE 5425", 3000 + i)
    for i in range(5):
        lectura("2026-08-%02d" % (i + 1), "TRACTOR MASSEY FERGUSON 6711", 2000 + i)
    lectura("2026-09-01", "CAMION", 104000)    # UNA sola, pero la mas reciente
    ws.append(["2026-09-02", "10:00", "LABOR", "Poda", "NOGALES", "", 1, "", "",
               None, "", "t", "Juan Parada", "", None, None, None, None])
    ruta = tmp_path / "maquinas.xlsx"
    wb.save(ruta)
    return str(ruta)


def test_manda_cuantas_veces_se_leyo_no_cuando_fue_la_ultima(tmp_path):
    """El caso real: CAMION tiene UNA lectura del 10-ago y le ganaba el boton
    a un tractor con 10 lecturas, por un dia de diferencia."""
    m = maquinas_para_botones(_excel_maquinas(tmp_path))
    assert m[0] == "TRACTOR JOHN DEERE 5425"
    assert m.index("TRACTOR MASSEY FERGUSON 6711") < m.index("CAMION")


def test_una_fila_sin_odometro_no_cuenta_como_lectura(tmp_path):
    """La fila de LABOR no lleva maquina ni odometro: no puede sumar."""
    m = maquinas_para_botones(_excel_maquinas(tmp_path))
    assert "" not in m
    assert len(m) == 3


def test_no_devuelve_una_pared_de_botones(tmp_path):
    m = maquinas_para_botones(_excel_maquinas(tmp_path), tope=2)
    assert len(m) == 2


def test_sin_hoja_bitacora_devuelve_lista_vacia(tmp_path):
    wb = Workbook()
    ruta = tmp_path / "vacio.xlsx"
    wb.save(ruta)
    assert maquinas_para_botones(str(ruta)) == []


def test_un_excel_ilegible_no_revienta_las_maquinas(tmp_path):
    assert maquinas_para_botones(str(tmp_path / "no_existe.xlsx")) == []


def _excel(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws.append(["Fecha", "Hora", "Tipo", "Actividad"])
    for _ in range(9):
        ws.append(["2026-08-20", "10:00", "LABOR", "Poda nogales"])
    for _ in range(4):
        ws.append(["2026-08-20", "10:00", "LABOR", "Aplicación herbicida"])
    for _ in range(7):
        ws.append(["2026-08-20", "10:00", "MAQUINARIA", "Lectura de horómetro"])
    ws.append(["2026-08-20", "10:00", "LABOR", "poda nogales"])   # misma, otra grafía
    ruta = tmp_path / "master.xlsx"
    wb.save(ruta)
    return str(ruta)


def test_las_labores_van_de_mas_a_menos_usada(tmp_path):
    l = labores_frecuentes(_excel(tmp_path))
    assert l[0] == "Poda nogales"
    assert "Aplicación herbicida" in l


def test_lectura_de_horometro_no_es_una_labor(tmp_path):
    """La escribe el propio bot; ofrecersela a Juan como labor no tiene sentido."""
    assert "Lectura de horómetro" not in labores_frecuentes(_excel(tmp_path))


def test_la_misma_labor_con_otra_grafia_no_se_cuenta_dos_veces(tmp_path):
    l = labores_frecuentes(_excel(tmp_path))
    assert sum(1 for x in l if x.lower() == "poda nogales") == 1


def test_un_excel_ilegible_no_revienta(tmp_path):
    assert labores_frecuentes(str(tmp_path / "no_existe.xlsx")) == []


def _excel_fechas_mezcladas(tmp_path):
    """La columna Fecha guarda LAS DOS formas: datetime en la mayoria de las
    filas y la cadena '2026-05-05' en unas 178 del Master real. Una fecha
    escrita a mano en dd/mm/aaaa entra por el mismo agujero."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws.append(["Fecha", "Hora", "Tipo", "Actividad", "Cultivo", "Sector",
               "Jornadas Hombre", "Trabajadores", "Insumo", "Cantidad",
               "Unidad", "Registro", "Registrado por", "Máquina", "Odómetro",
               "Horas Día", "Superficie ha", "Días Cubiertos"])

    def lectura(fecha, maquina, odo):
        ws.append([fecha, "10:00", "MAQUINARIA", "Lectura de horómetro",
                   "GENERAL", "", None, "", "", None, "", "t", "Juan Parada",
                   maquina, odo, None, None, None])

    # Dos maquinas con EL MISMO numero de lecturas: manda el desempate.
    lectura("2026-04-01", "TRACTOR RECIENTE", 100)
    lectura("05/09/2026", "TRACTOR RECIENTE", 110)     # 5-sep, escrita a mano
    lectura("2026-04-01", "TRACTOR ANTIGUO", 200)
    lectura("2026-05-05", "TRACTOR ANTIGUO", 210)      # 5-may, cadena ISO
    ruta = tmp_path / "fechas.xlsx"
    wb.save(ruta)
    return str(ruta)


def test_una_fecha_escrita_a_mano_no_desordena_el_desempate(tmp_path):
    """Comparadas como TEXTO, '2026-05-05' > '05/09/2026' y el tractor de mayo
    le gana el boton al de septiembre. Hay que comparar fechas, no cadenas."""
    m = maquinas_para_botones(_excel_fechas_mezcladas(tmp_path))
    assert m.index("TRACTOR RECIENTE") < m.index("TRACTOR ANTIGUO")


def test_la_misma_labor_con_doble_espacio_o_sin_tilde_es_una_sola(tmp_path):
    """`.lower()` solo no basta: 'Poda  Nogales' y 'Aplicacion herbicida' se
    cuentan aparte y se comen dos botones de los seis."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws.append(["Fecha", "Hora", "Tipo", "Actividad"])
    for _ in range(5):
        ws.append(["2026-08-20", "10:00", "LABOR", "Poda nogales"])
    ws.append(["2026-08-20", "10:00", "LABOR", "Poda  nogales"])   # doble espacio
    for _ in range(3):
        ws.append(["2026-08-20", "10:00", "LABOR", "Aplicación herbicida"])
    ws.append(["2026-08-20", "10:00", "LABOR", "Aplicacion herbicida"])  # sin tilde
    ruta = tmp_path / "labores.xlsx"
    wb.save(ruta)

    l = labores_frecuentes(str(ruta))
    assert l == ["Poda nogales", "Aplicación herbicida"]
