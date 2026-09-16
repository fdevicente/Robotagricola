# -*- coding: utf-8 -*-
"""Conciliar una cuota no es pagar la factura.

Revisión del 14-sep-2026 (informe "¿Llevar el Robot a una Raspberry Pi?"):

  · `registrar_vinculos` escribe la Fecha Pago de la factura **siempre**, aunque
    el movimiento sea una cuota. Con eso la factura sale de la deuda: el
    dashboard cuenta "pagada" cualquier fila con fecha. Y el modelo justamente
    existe para pagos parciales (S-Invest, plantas de avellano).
  · `desconciliar` borra el vínculo pero **deja la fecha puesta**: la deuda no
    vuelve, y nadie se entera.
  · Nada impide vincular **dos veces el mismo pago al mismo documento**: lo
    asignado se duplicaría y la factura quedaría "pagada" con la mitad.

La cobertura ya se sabe calcular (`estado_documento`: pagado con saldo ≤ $1,
si no parcial). Lo que falta es usarla antes de escribir la fecha.
"""
import pytest
from openpyxl import Workbook, load_workbook

from modules.conciliacion_store import (HEADERS, SHEET, asignado_por_documento,
                                        desconciliar, registrar_vinculos)

BANCO = ["Fecha", "Descripcion", "Referencia", "Cargo", "Abono", "Saldo",
         "Tipo", "Categoria", "Cultivo", "Factura_linkeada"]
FACTURAS = ["Fecha Emision", "Fecha Vencimiento", "Fecha Pago", "Proveedor",
            "Rut", "Documento", "Numero", "Glosa", "Glosa II", "Valor unitario",
            "Cantidad", "TOTAL NETO", "IVA", "Impuesto", "Monto / TOTAL",
            "Total Factura"]


@pytest.fixture
def libro(tmp_path):
    """Master de prueba con Facturas, Cuenta Banco y Conciliaciones."""
    def _crear(facturas, movimientos, vinculos=()):
        wb = Workbook()
        ws = wb.active
        ws.title = "Facturas"
        ws.append(FACTURAS)
        for proveedor, nro, monto, fecha_pago in facturas:
            fila = [None] * 16
            fila[2], fila[3], fila[6] = fecha_pago, proveedor, nro
            fila[14] = monto
            ws.append(fila)
        # El Total Factura (col 16) es el ancla del documento: el MISMO en
        # todas sus líneas, no el monto del ítem.
        totales = {}
        for f in range(2, ws.max_row + 1):
            clave = (ws.cell(f, 4).value, str(ws.cell(f, 7).value))
            totales[clave] = totales.get(clave, 0) + float(ws.cell(f, 15).value or 0)
        for f in range(2, ws.max_row + 1):
            clave = (ws.cell(f, 4).value, str(ws.cell(f, 7).value))
            ws.cell(f, 16).value = totales[clave]
        wb_banco = wb.create_sheet("Cuenta Banco")
        wb_banco.append(BANCO)
        for fecha, desc, cargo in movimientos:
            wb_banco.append([fecha, desc, "", cargo, None, None, "", "", "", ""])
        wc = wb.create_sheet(SHEET)
        wc.append(HEADERS)
        for i, (fila_banco, nro, prov, asignado) in enumerate(vinculos, 1):
            wc.append([i, "2026-08-01", fila_banco, "2026-07-01", "mov", 0,
                       "FACTURA", None, nro, prov, asignado, "manual", "", ""])
        ruta = tmp_path / "master.xlsx"
        wb.save(ruta)
        wb.close()
        return str(ruta)
    return _crear


def _fechas_pago(ruta):
    wb = load_workbook(ruta)
    ws = wb["Facturas"]
    fechas = [ws.cell(f, 3).value for f in range(2, ws.max_row + 1)]
    wb.close()
    return fechas


def _vinculo(fila_banco, nro, proveedor, filas_doc, monto=None, fecha_pago=None):
    return {"fila_banco": fila_banco, "tipo_doc": "FACTURA",
            "fila_doc": filas_doc[0], "filas_doc": filas_doc,
            "nro_doc": nro, "proveedor": proveedor, "monto_asignado": monto,
            "criterio": "manual", "fecha_pago": fecha_pago}


# ── La fecha de pago solo cuando el pago cubre la factura ──────────────────

def test_una_cuota_no_pone_fecha_de_pago(libro):
    """S-Invest: 5 de 10 millones. La factura sigue debiendo la mitad."""
    p = libro(facturas=[("S-INVEST", "1928", 10_000_000, None)],
              movimientos=[("2026-05-12", "S-Invest cuota 1", 5_000_000)])

    registrar_vinculos([_vinculo(2, "1928", "S-INVEST", [2], 5_000_000)],
                       usuario="test", excel_path=p)

    assert _fechas_pago(p) == [None], "marcó pagada una factura pagada a medias"


def test_la_cuota_que_completa_el_total_si_pone_la_fecha(libro):
    """La segunda cuota cierra la factura: ahí sí queda pagada."""
    p = libro(facturas=[("S-INVEST", "1928", 10_000_000, None)],
              movimientos=[("2026-05-12", "cuota 1", 5_000_000),
                           ("2026-06-03", "cuota 2", 5_000_000)],
              vinculos=[(2, "1928", "S-INVEST", 5_000_000)])

    registrar_vinculos([_vinculo(3, "1928", "S-INVEST", [2], 5_000_000)],
                       usuario="test", excel_path=p)

    assert _fechas_pago(p) == ["2026-06-03"]


def test_un_pago_que_cubre_todo_pone_la_fecha(libro):
    """Lo de siempre: un cargo que paga la factura entera la deja pagada."""
    p = libro(facturas=[("COPEVAL", "6231249", 307_199, None)],
              movimientos=[("2026-04-23", "Copeval", 307_199)])

    registrar_vinculos([_vinculo(2, "6231249", "COPEVAL", [2], 307_199)],
                       usuario="test", excel_path=p)

    assert _fechas_pago(p) == ["2026-04-23"]


def test_todas_las_lineas_de_la_factura_quedan_con_la_fecha(libro):
    """Una factura de tres ítems son tres filas: o todas o ninguna."""
    p = libro(facturas=[("COPEVAL", "77", 1000, None),
                        ("COPEVAL", "77", 2000, None),
                        ("COPEVAL", "77", 3000, None)],
              movimientos=[("2026-04-23", "Copeval", 6000)])

    registrar_vinculos([_vinculo(2, "77", "COPEVAL", [2, 3, 4], 6000)],
                       usuario="test", excel_path=p)

    assert _fechas_pago(p) == ["2026-04-23"] * 3


# ── Deshacer la conciliación devuelve la deuda ─────────────────────────────

def test_desconciliar_borra_la_fecha_que_puso_la_conciliacion(libro):
    """Si el vínculo se va, la factura vuelve a deber."""
    p = libro(facturas=[("COPEVAL", "6231249", 307_199, "2026-04-23")],
              movimientos=[("2026-04-23", "Copeval", 307_199)],
              vinculos=[(2, "6231249", "COPEVAL", 307_199)])

    assert desconciliar(1, excel_path=p) is True

    assert _fechas_pago(p) == [None]
    assert asignado_por_documento(p) == {}


def test_desconciliar_no_borra_una_fecha_puesta_a_mano(libro):
    """El dueño la escribió con otra fecha: esa no se toca."""
    p = libro(facturas=[("COPEVAL", "6231249", 307_199, "2026-01-05")],
              movimientos=[("2026-04-23", "Copeval", 307_199)],
              vinculos=[(2, "6231249", "COPEVAL", 307_199)])

    desconciliar(1, excel_path=p)

    assert _fechas_pago(p) == ["2026-01-05"]


# ── El mismo pago no se vincula dos veces ──────────────────────────────────

def test_no_se_vincula_dos_veces_el_mismo_pago_al_mismo_documento(libro):
    """Vincular dos veces duplicaría lo asignado y la dejaría 'pagada' con la mitad."""
    p = libro(facturas=[("S-INVEST", "1928", 10_000_000, None)],
              movimientos=[("2026-05-12", "cuota 1", 5_000_000)],
              vinculos=[(2, "1928", "S-INVEST", 5_000_000)])

    r = registrar_vinculos([_vinculo(2, "1928", "S-INVEST", [2], 5_000_000)],
                           usuario="test", excel_path=p)

    assert r["registrados"] == 0
    assert r["repetidos"] == 1
    assert list(asignado_por_documento(p).values()) == [5_000_000]
    assert _fechas_pago(p) == [None]
