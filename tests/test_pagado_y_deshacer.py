# -*- coding: utf-8 -*-
"""/pagado y /deshacer no pueden tocar filas que no son.

Revisión del 14-sep-2026 (informe "¿Llevar el Robot a una Raspberry Pi?"):

  · **/pagado** escribe la fecha de pago en TODAS las filas con ese número, de
    cualquier proveedor, y pisa la fecha que ya estaba (`registrar_pago`). El
    propio repo tiene medido que **19 números de factura los usan dos o tres
    proveedores distintos** (`modules/drive/enlaces.py`): el Nº1034 es de
    AGROAVELLANO y de LUCIA GARCIA, el Nº264 de ECOSMART y de INV. SANTA
    VICTORIA. Dar por pagada la factura de otro la saca de la deuda sin que
    nadie lo note.
  · **/deshacer** borra las últimas N filas sin mirar qué son
    (`delete_last_rows`). Si Juan guardó una factura después de la tuya, borra
    la de Juan. Con boletas, además borra la última fila de Caja Chica, sea lo
    que sea.

Lo que se espera: /pagado pregunta cuando el número está repetido y no pisa
fechas; /deshacer comprueba que lo último del Excel sea lo que guardó quien
escribe, y si no, no borra nada.
"""
import asyncio
import types

import openpyxl
import pytest

import excel_manager

CAJA_HEADERS = ["Fecha", "Tipo", "Detalle", "Comercio / Proveedor", "Nº Boleta",
                "Ingreso", "Egreso", "Saldo"]


def _master(ruta, facturas=(), boletas=(), caja=()):
    """Master mínimo: Facturas (proveedor, nº, monto, fecha_pago) y, si se
    piden, Boletas (proveedor, nº, monto) y Caja Chica (proveedor, nº, egreso)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.append(["Fecha Emision", "Fecha Vencimiento", "Fecha Pago", "Proveedor",
               "Rut", "Documento", "Numero", "Glosa", "Glosa II", "Valor unitario",
               "Cantidad", "TOTAL NETO", "IVA", "Impuesto", "Monto / TOTAL",
               "Total Factura"])
    for proveedor, numero, monto, fecha_pago in facturas:
        fila = [None] * 16
        fila[2], fila[3], fila[6] = fecha_pago, proveedor, numero
        fila[14] = fila[15] = monto
        ws.append(fila)
    if boletas or caja:
        ws_b = wb.create_sheet("Boletas")
        ws_b.append(["Fecha", "Proveedor", "Rut", "Detalle", "Monto", "Nº Doc", "Glosa II"])
        for proveedor, numero, monto in boletas:
            ws_b.append([None, proveedor, None, "compra", monto, numero, None])
        ws_c = wb.create_sheet("Caja Chica")
        ws_c.append(CAJA_HEADERS)
        for proveedor, numero, egreso in caja:
            ws_c.append([None, "Gasto", "compra", proveedor, numero, None, egreso, 0])
    wb.save(ruta)
    wb.close()


def _facturas(ruta):
    """(proveedor, número, fecha de pago) de cada fila."""
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Facturas"]
    filas = [(ws.cell(f, 4).value, ws.cell(f, 7).value, ws.cell(f, 3).value)
             for f in range(2, ws.max_row + 1)]
    wb.close()
    return filas


def _hoja(ruta, nombre, columnas):
    wb = openpyxl.load_workbook(ruta)
    ws = wb[nombre]
    filas = [tuple(ws.cell(f, c).value for c in columnas)
             for f in range(2, ws.max_row + 1)]
    wb.close()
    return filas


@pytest.fixture
def master(tmp_path, monkeypatch):
    """El Master de estos tests, apuntado por el módulo."""
    ruta = tmp_path / "MASTER.xlsx"
    monkeypatch.setattr(excel_manager, "EXCEL_PATH", str(ruta))
    return ruta


# ── /pagado ────────────────────────────────────────────────────────────────

def test_pagar_no_toca_la_factura_del_otro_proveedor(master):
    """El Nº1034 es de AGROAVELLANO y de LUCIA GARCIA."""
    _master(master, facturas=[("AGROAVELLANO", "1034", 100000, None),
                              ("LUCIA GARCIA", "1034", 50000, None)])

    excel_manager.registrar_pago("1034", "2026-09-15", "Banco",
                                 proveedor="AGROAVELLANO")

    assert _facturas(master) == [("AGROAVELLANO", "1034", "2026-09-15 (Banco)"),
                                 ("LUCIA GARCIA", "1034", None)]


def test_con_el_numero_repetido_y_sin_proveedor_no_escribe_nada(master):
    """Sin saber de quién es, marcarla es peor que no hacer nada."""
    _master(master, facturas=[("AGROAVELLANO", "1034", 100000, None),
                              ("LUCIA GARCIA", "1034", 50000, None)])

    r = excel_manager.registrar_pago("1034", "2026-09-15", "Banco")

    assert r["actualizadas"] == 0
    assert sorted(r["proveedores"]) == ["AGROAVELLANO", "LUCIA GARCIA"]
    assert [f[2] for f in _facturas(master)] == [None, None]


def test_no_pisa_la_fecha_de_pago_que_ya_estaba(master):
    """Una fecha puesta a mano vale más que una nueva: se avisa, no se pisa."""
    _master(master, facturas=[("COPEVAL", "55", 1000, "2026-06-01")])

    r = excel_manager.registrar_pago("55", "2026-09-15", "Banco",
                                     proveedor="COPEVAL")

    assert _facturas(master) == [("COPEVAL", "55", "2026-06-01")]
    assert r["ya_tenian"] == 1


def test_con_un_solo_proveedor_marca_todas_sus_lineas(master):
    """Lo que ya funcionaba: una factura de tres ítems se paga entera."""
    _master(master, facturas=[("COPEVAL", "77", 1000, None),
                              ("COPEVAL", "77", 2000, None),
                              ("COPEVAL", "77", 3000, None)])

    r = excel_manager.registrar_pago("77", "2026-09-15", "Caja Chica")

    assert r["actualizadas"] == 3
    assert all(f[2] for f in _facturas(master))


# ── /deshacer ──────────────────────────────────────────────────────────────

def test_deshacer_no_borra_la_factura_que_guardo_otro(master):
    """Juan guardó la suya después: deshacer la mía ya no es borrar la última."""
    _master(master, facturas=[("PROVEEDOR A", "101", 1000, None),
                              ("PROVEEDOR B", "202", 2000, None)])

    ok = excel_manager.delete_last_rows(1, identidad=("PROVEEDOR A", "101"))

    assert ok is False
    assert _facturas(master) == [("PROVEEDOR A", "101", None),
                                 ("PROVEEDOR B", "202", None)]


def test_deshacer_borra_lo_suyo_cuando_calza(master):
    _master(master, facturas=[("PROVEEDOR B", "202", 2000, None),
                              ("PROVEEDOR A", "101", 1000, None),
                              ("PROVEEDOR A", "101", 500, None)])

    ok = excel_manager.delete_last_rows(2, identidad=("PROVEEDOR A", "101"))

    assert ok is True
    assert _facturas(master) == [("PROVEEDOR B", "202", None)]


def test_deshacer_una_boleta_no_borra_el_gasto_de_caja_chica_ajeno(master):
    """La fila de Caja Chica que sobra es de otra boleta."""
    _master(master,
            boletas=[("PROVEEDOR A", "B-1", 5000), ("PROVEEDOR B", "B-2", 9000)],
            caja=[("PROVEEDOR A", "B-1", 5000), ("PROVEEDOR B", "B-2", 9000)])

    ok = excel_manager.delete_last_boletas(1, identidad=("PROVEEDOR A", "B-1"))

    assert ok is False
    assert _hoja(master, "Boletas", (2, 6)) == [("PROVEEDOR A", "B-1"),
                                                ("PROVEEDOR B", "B-2")]
    assert _hoja(master, "Caja Chica", (4, 5)) == [("PROVEEDOR A", "B-1"),
                                                   ("PROVEEDOR B", "B-2")]


def test_deshacer_una_boleta_borra_su_fila_y_su_gasto(master):
    _master(master,
            boletas=[("PROVEEDOR B", "B-2", 9000), ("PROVEEDOR A", "B-1", 5000)],
            caja=[("PROVEEDOR B", "B-2", 9000), ("PROVEEDOR A", "B-1", 5000)])

    ok = excel_manager.delete_last_boletas(1, identidad=("PROVEEDOR A", "B-1"))

    assert ok is True
    assert _hoja(master, "Boletas", (2, 6)) == [("PROVEEDOR B", "B-2")]
    assert _hoja(master, "Caja Chica", (4, 5)) == [("PROVEEDOR B", "B-2")]


# ── En el flujo del bot ────────────────────────────────────────────────────

class _Msg:
    def __init__(self, texto=""):
        self.text = texto
        self.respuestas = []

    async def reply_text(self, texto, **kw):
        self.respuestas.append(texto)
        return _Msg()

    async def edit_text(self, texto, **kw):
        self.respuestas.append(texto)


def test_pagado_con_numero_repetido_pregunta_de_quien_es(monkeypatch):
    """Antes mostraba el primer proveedor de la lista y seguía como si nada."""
    from handlers import finanzas
    monkeypatch.setattr(finanzas, "buscar_factura", lambda nro: [
        {"fila": 5, "proveedor": "AGROAVELLANO", "rut": "1-9", "documento": "Factura",
         "nro_factura": "1034", "glosa": "x", "total": 100000, "fecha_pago": None},
        {"fila": 9, "proveedor": "LUCIA GARCIA", "rut": "2-7", "documento": "Factura",
         "nro_factura": "1034", "glosa": "y", "total": 50000, "fecha_pago": None},
    ])
    mensaje = _Msg("1034")
    update = types.SimpleNamespace(message=mensaje)
    ctx = types.SimpleNamespace(user_data={"pagado_state": "esperando_nro"})

    asyncio.run(finanzas.handle_text_pagado(update, ctx))

    assert ctx.user_data["pagado_state"] != "esperando_fecha", \
        "siguió adelante sin saber de qué proveedor es la factura"
    texto = " ".join(mensaje.respuestas)
    assert "AGROAVELLANO" in texto and "LUCIA GARCIA" in texto


def test_pagado_acepta_el_proveedor_escrito_a_mano(monkeypatch):
    """Los botones no pueden ser la única puerta: un estado que no escucha lo
    que le escriben es la misma trampa del `/ cancelar` y del horómetro."""
    from handlers import finanzas
    ctx = types.SimpleNamespace(user_data={
        "pagado_state": "esperando_proveedor", "pagado_nro": "1034",
        "pagado_proveedores": ["AGROAVELLANO", "LUCIA GARCIA"]})
    mensaje = _Msg("lucia garcia")
    update = types.SimpleNamespace(message=mensaje)

    asyncio.run(finanzas.handle_text_pagado(update, ctx))

    assert ctx.user_data["pagado_proveedor"] == "LUCIA GARCIA"
    assert ctx.user_data["pagado_state"] == "esperando_fecha"


def test_deshacer_explica_por_que_no_borro(monkeypatch):
    """Cuando el Excel se niega, el dueño tiene que entender qué pasó: decirle
    solo "error" es esconder que el bot hizo lo correcto."""
    import main

    class _MsgEncadenado(_Msg):
        async def reply_text(self, texto, **kw):
            self.respuestas.append(texto)
            return self          # para poder leer lo que se edita después

    monkeypatch.setattr(main, "delete_last_rows", lambda n, identidad=None: False)
    mensaje = _MsgEncadenado("/deshacer")
    update = types.SimpleNamespace(message=mensaje)
    ctx = types.SimpleNamespace(user_data={
        "last_invoice_file": "no_existe.jpg", "last_invoice_rows": 1,
        "last_invoice_boleta": False, "last_invoice_id": ("PROVEEDOR A", "101")})

    asyncio.run(main.cmd_deshacer(update, ctx))

    texto = " ".join(mensaje.respuestas).lower()
    assert "no borré nada" in texto, mensaje.respuestas


def test_deshacer_le_dice_al_excel_que_factura_era(monkeypatch):
    """Sin la identidad, el Excel no puede negarse a borrar lo ajeno."""
    import main

    recibido = {}

    def falso_delete(n, identidad=None):
        recibido["n"], recibido["identidad"] = n, identidad
        return True

    monkeypatch.setattr(main, "delete_last_rows", falso_delete)
    mensaje = _Msg("/deshacer")
    update = types.SimpleNamespace(message=mensaje)
    ctx = types.SimpleNamespace(user_data={
        "last_invoice_file": "no_existe.jpg", "last_invoice_rows": 2,
        "last_invoice_boleta": False,
        "last_invoice_id": ("PROVEEDOR A", "101")})

    asyncio.run(main.cmd_deshacer(update, ctx))

    assert recibido["n"] == 2
    assert recibido["identidad"] == ("PROVEEDOR A", "101")
