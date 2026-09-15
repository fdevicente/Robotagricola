# -*- coding: utf-8 -*-
"""Las vacaciones se RECALCULAN, no se suman a ciegas.

🔴 Medido el 15-sep-2026:
  · actualizar_dias_mensuales sumaba +1,25 cada vez que corría, sin anotar el
    mes. El 1-jul falló (Master bloqueado) y solo lo dejó en el log; el 1-ago y
    el 1-sep el bot estaba abajo porque el PC dormía. La acumulación mensual
    NUNCA se hizo con éxito desde el bot y hubo que corregir a mano.
  · registrar_vacacion contaba hábiles con toordinal() % 7: saltaba viernes y
    sábado y contaba el domingo. Una semana de lunes a viernes daba 4.

Ahora el saldo sale de la misma cuenta que usó la carga del 4-ago: saldo base
(hoja Vacaciones Pendientes) + 1,25 por mes − días tomados (hoja Vacaciones).
Correrla dos veces, o tarde, da lo mismo.
"""
import asyncio
import inspect
from datetime import date

import openpyxl
import pytest

import vacaciones_manager as vm

SEPTIEMBRE = date(2026, 9, 1)


@pytest.fixture
def master(tmp_path):
    ruta = tmp_path / "master.xlsx"
    wb = openpyxl.Workbook()
    per = wb.active
    per.title = "Personal"
    per.append(vm.PERSONAL_HEADERS + ["Notas"])
    per.append(["Ana Uno", "1-9", "Operaria", date(2025, 1, 1), 0, 0, None, ""])
    per.append(["Beto Dos", "2-7", "Operario", date(2026, 1, 1), 0, 0, None, ""])
    vac = wb.create_sheet("Vacaciones")
    vac.append(vm.VACACIONES_HEADERS)
    vac.append(["Ana Uno", date(2026, 3, 2), date(2026, 3, 6), 5, "Aprobado", ""])
    base = wb.create_sheet("Vacaciones Pendientes")
    base.append(["Nombre", "RUT", "Fecha Contrato", "Saldo Último Conocido (días)",
                 "Fecha del Saldo", "Notas"])
    base.append(["Ana Uno", "1-9", date(2025, 1, 1), 10.0, date(2026, 1, 15), ""])
    wb.save(ruta)
    return str(ruta)


def _personal(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    out = {r[0]: {"pendientes": r[4], "tomados": r[5], "ultima": r[6]}
           for r in wb["Personal"].iter_rows(min_row=2, values_only=True) if r[0]}
    wb.close()
    return out


# ── El saldo se recalcula ───────────────────────────────────────────────────

def test_el_saldo_es_la_base_mas_lo_acumulado_menos_lo_tomado(master):
    vm.actualizar_dias_mensuales(hasta=SEPTIEMBRE, path=master)
    p = _personal(master)
    assert p["Ana Uno"]["pendientes"] == 15.0      # 10 + 8 meses × 1,25 − 5
    assert p["Ana Uno"]["tomados"] == 5
    assert p["Beto Dos"]["pendientes"] == 10.0     # sin base: parte de su ingreso


def test_correr_dos_veces_el_mismo_mes_no_duplica(master):
    vm.actualizar_dias_mensuales(hasta=SEPTIEMBRE, path=master)
    vm.actualizar_dias_mensuales(hasta=SEPTIEMBRE, path=master)
    assert _personal(master)["Ana Uno"]["pendientes"] == 15.0


def test_si_un_mes_no_corrio_el_siguiente_se_pone_al_dia(master):
    vm.actualizar_dias_mensuales(hasta=date(2026, 8, 1), path=master)
    assert _personal(master)["Ana Uno"]["pendientes"] == 13.75
    vm.actualizar_dias_mensuales(hasta=date(2026, 10, 1), path=master)   # septiembre no corrió
    assert _personal(master)["Ana Uno"]["pendientes"] == 16.25


def test_si_nada_cambio_no_vuelve_a_guardar_el_master(master, monkeypatch):
    vm.actualizar_dias_mensuales(hasta=SEPTIEMBRE, path=master)
    guardados = []
    monkeypatch.setattr(vm, "_save_wb", lambda wb, path=None: guardados.append(path))
    r = vm.actualizar_dias_mensuales(hasta=SEPTIEMBRE, path=master)
    assert r["actualizados"] == 0
    assert guardados == []


# ── Días hábiles ────────────────────────────────────────────────────────────

def test_una_semana_de_lunes_a_viernes_son_5_dias_habiles():
    assert vm.dias_habiles(date(2026, 9, 21), date(2026, 9, 25)) == 5


def test_el_fin_de_semana_no_cuenta():
    assert vm.dias_habiles(date(2026, 9, 25), date(2026, 9, 28)) == 2   # vie, sáb, dom, lun


def test_los_feriados_no_cuentan():
    """El 18-sep-2026 es viernes y feriado: esa semana tiene 4 hábiles."""
    assert vm.dias_habiles(date(2026, 9, 14), date(2026, 9, 18)) == 4


def test_registrar_vacacion_cuenta_bien_y_deja_el_saldo_recalculado(master):
    r = vm.registrar_vacacion("Beto Dos", "2026-09-21", "2026-09-25",
                              path=master, hasta=SEPTIEMBRE)
    assert r["dias_habiles"] == 5
    p = _personal(master)["Beto Dos"]
    assert p["tomados"] == 5
    assert p["pendientes"] == 5.0                  # 8 meses × 1,25 − 5
    assert p["ultima"].date() == date(2026, 9, 25)


# ── El job del bot ──────────────────────────────────────────────────────────

class _Bot:
    def __init__(self):
        self.enviados = []

    async def send_message(self, chat_id, text, **kw):
        self.enviados.append((chat_id, text))


class _Ctx:
    def __init__(self):
        self.bot = _Bot()
        self.bot_data = {"owner_chat_id": 8684368429}


def test_si_la_actualizacion_falla_le_avisa_al_dueno(monkeypatch):
    """El 1-jul falló por el Master bloqueado y solo quedó en el log."""
    from handlers import personal

    def falla(*a, **k):
        raise PermissionError("[Errno 13] Permission denied: MASTER")

    monkeypatch.setattr(personal, "actualizar_dias_mensuales", falla)
    ctx = _Ctx()
    asyncio.run(personal.job_vacaciones_mensuales(ctx))
    assert len(ctx.bot.enviados) == 1
    assert "vacaciones" in ctx.bot.enviados[0][1].lower()


def test_si_no_cambio_nada_no_manda_mensaje(monkeypatch):
    """Como corre seguido, avisar cada vez sería ruido."""
    from handlers import personal
    monkeypatch.setattr(personal, "actualizar_dias_mensuales",
                        lambda *a, **k: {"actualizados": 0, "incremento": 1.25,
                                         "hasta": "2026-09"})
    ctx = _Ctx()
    asyncio.run(personal.job_vacaciones_mensuales(ctx))
    assert ctx.bot.enviados == []


def test_el_bot_revisa_las_vacaciones_seguido_y_no_solo_el_dia_1():
    """Con run_monthly, si el PC dormía el día 1 a las 07:00, ese mes se perdía."""
    import main
    assert "run_repeating(job_vacaciones_mensuales" in inspect.getsource(main)
