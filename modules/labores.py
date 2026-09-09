# -*- coding: utf-8 -*-
"""Cuanto trabajo se hizo y cuanto costo.

Modulo PURO: lee el Master, no escribe nada y no sabe de Flask. Asi se prueba
entero en memoria y el mismo dato sirve despues por Telegram.

DE DONDE SALE EL COSTO
No hay ningun valor de jornada en el Master --la hoja Personal no tiene
sueldos-- y pedirle al dueño que cargue una tarifa envejece en silencio. Lo que
si hay es plata que salio: las transferencias a la gente de planta llevan el RUT
en la glosa ("TEF 9850887-2 FELICITO AMIGO") y esos RUT estan en Personal. El
cruce es por RUT, no por nombre.
"""
import datetime as dt
import logging
import re
import unicodedata

logger = logging.getLogger(__name__)

BITACORA_SHEET = "Bitácora"

# La escribe el propio bot al guardar una lectura: no es trabajo de nadie.
_NO_ES_LABOR = {"lectura de horometro"}

# NO ES TRABAJO, aunque ocupe una fila de la bitacora. Medido contra el Master:
# por Tipo hay LABOR 313 jornadas, RIEGO 41, APLICACION 24, MAQUINARIA 4 y
# OTRO 0. Sin este filtro, "Vacaciones" y "Ausente" salian en la lista de
# labores como si alguien hubiera trabajado.
_TIPO_NO_TRABAJO = {"OTRO"}
_AUSENCIAS = {"vacaciones", "ausente", "aucente", "sin uso", "licencia"}

# Agrupacion por palabra clave. EL ORDEN IMPORTA: gana la primera que calza y
# las mas especificas van primero. "Sacar restos poda nogales" tiene que caer en
# su grupo ANTES de que "poda" se lo lleve; son trabajos distintos y medido
# contra el Master, sacar los restos costo mas del doble de jornadas que podar.
GRUPOS = [
    ("sacar restos", "Sacar restos de poda"),
    ("pintar", "Pintar poda"),
    ("bajar ramas", "Bajar ramas"),
    ("poda", "Poda"),
    ("herbicida", "Aplicación herbicida"),
    ("fungicida", "Aplicación fungicida"),
    ("fertiliza", "Fertilización"),
    ("desagu", "Desaguar"),
    ("riego", "Mantención riego"),
    ("maquinaria", "Mantención maquinaria"),
    ("planta", "Mantención planta"),
    ("mantencion", "Mantención general"),
    ("aseo", "Aseo"),
    ("rastra", "Pasar rastra"),
    ("replante", "Replante"),
    ("cosecha", "Cosecha"),
]


def _clave(t) -> str:
    """Minusculas, sin tildes y con los espacios colapsados."""
    t = "".join(c for c in unicodedata.normalize("NFD", str(t or "").lower())
                if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", t).split())


def grupo_de(actividad) -> str:
    """A que labor pertenece una etiqueta. La etiqueta cruda si no calza."""
    k = _clave(actividad)
    for palabra, nombre in GRUPOS:
        if palabra in k:
            return nombre
    return str(actividad or "").strip()


def _fecha(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    for molde in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(str(v or "")[:10], molde).date()
        except ValueError:
            continue
    return None


def _leer_bitacora(path=None) -> list:
    """Filas utiles de la bitacora, ya normalizadas."""
    from openpyxl import load_workbook

    from config import EXCEL_PATH
    filas = []
    try:
        wb = load_workbook(path or EXCEL_PATH, read_only=True, data_only=True)
        try:
            if BITACORA_SHEET not in wb.sheetnames:
                return []
            ws = wb[BITACORA_SHEET]
            cab = next(ws.iter_rows(min_row=1, max_row=1), None)
            enc = [c.value for c in cab] if cab else []
            idx = {n: i for i, n in enumerate(enc)}

            def col(r, n):
                return r[idx[n]] if n in idx and len(r) > idx[n] else None

            for r in ws.iter_rows(min_row=2, values_only=True):
                if not r or not any(r):
                    continue
                act = str(col(r, "Actividad") or "").strip()
                if not act or _clave(act) in _NO_ES_LABOR:
                    continue
                if _clave(act) in _AUSENCIAS:
                    continue
                tipo = str(col(r, "Tipo") or "").strip().upper()
                if tipo in _TIPO_NO_TRABAJO:
                    continue
                try:
                    jh = float(col(r, "Jornadas Hombre") or 0)
                except (TypeError, ValueError):
                    jh = 0.0
                # Una fila de MAQUINARIA sin jornadas es el bot anotando la
                # maquina --combustible, horometro roto--, no trabajo de nadie.
                # Con jornadas si lo es: "Destroncar nogales" entra por ahi.
                if tipo == "MAQUINARIA" and jh <= 0:
                    continue
                personas = [p.strip() for p in
                            str(col(r, "Trabajadores") or "").split(",")
                            if p.strip()]
                filas.append({
                    "fecha": _fecha(col(r, "Fecha")),
                    "actividad": act,
                    "jornadas": jh,
                    "personas": personas,
                    "cultivo": str(col(r, "Cultivo") or "").strip(),
                    "sector": str(col(r, "Sector") or "").strip(),
                })
        finally:
            wb.close()
    except Exception as e:
        logger.warning("labores: no pude leer la bitácora: %r", e)
        return []
    return filas


def _en_rango(f, desde, hasta) -> bool:
    if f is None:
        return False
    if desde and f < _fecha(desde):
        return False
    if hasta and f > _fecha(hasta):
        return False
    return True


def resumen_labores(desde=None, hasta=None, path=None) -> list:
    """Por labor agrupada: jornadas, dias, personas y las etiquetas que la componen.

    Las etiquetas van en la salida a proposito: sin verlas, la agrupacion es una
    caja negra y nadie puede corregirla.
    """
    acc = {}
    for f in _leer_bitacora(path):
        if not _en_rango(f["fecha"], desde, hasta):
            continue
        g = grupo_de(f["actividad"])
        d = acc.setdefault(g, {"labor": g, "jornadas": 0.0, "_dias": set(),
                               "_pers": set(), "_etq": set()})
        d["jornadas"] += f["jornadas"]
        d["_dias"].add(f["fecha"])
        d["_pers"].update(_clave(p) for p in f["personas"])
        d["_etq"].add(f["actividad"])

    salida = []
    for d in acc.values():
        salida.append({
            "labor": d["labor"],
            "jornadas": d["jornadas"],
            "dias": len(d["_dias"]),
            "personas": len(d["_pers"]),
            "desde": str(min(d["_dias"])) if d["_dias"] else "",
            "hasta": str(max(d["_dias"])) if d["_dias"] else "",
            "etiquetas": sorted(d["_etq"]),
        })
    return sorted(salida, key=lambda x: -x["jornadas"])
