# -*- coding: utf-8 -*-
"""Que botones ofrecerle a Juan. Salen del Master, no van escritos a mano.

Las labores cambian con la temporada --hoy poda, en marzo cosecha-- y una lista
fija envejece sin que nadie lo note: Juan terminaria apretando "Otra..." siempre
y el boton dejaria de servir.
"""
import logging

logger = logging.getLogger(__name__)

MAX_MAQUINAS = 6
MAX_LABORES = 6
BITACORA_SHEET = "Bitácora"

# La escribe el propio bot al guardar una lectura: no es una labor que Juan haga.
_NO_ES_LABOR = {"lectura de horómetro", "lectura de horometro"}


def maquinas_para_botones(excel_path: str | None = None,
                          tope: int = MAX_MAQUINAS) -> list:
    """Maquinas ordenadas por CUANTAS veces se les leyo el horometro.

    No por fecha de la ultima lectura. Medido el 3-sep-2026 contra el Master
    real: CAMION y una camioneta tienen UNA sola lectura cada una, del 10 y el
    11-ago --pinta a carga inicial de fichas, no a uso-- y por un dia de
    diferencia le ganaban el boton al TRACTOR JOHN DEERE 5425, que Juan si usa.
    La recurrencia distingue lo que se usa de lo que se cargo una vez; la fecha
    no. Se desempata por la lectura mas reciente.
    """
    from collections import Counter

    from openpyxl import load_workbook

    from config import EXCEL_PATH
    ruta = excel_path or EXCEL_PATH
    cuenta, ultima = Counter(), {}
    try:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            if BITACORA_SHEET not in wb.sheetnames:
                return []
            ws = wb[BITACORA_SHEET]
            fila_enc = next(ws.iter_rows(min_row=1, max_row=1), None)
            enc = [c.value for c in fila_enc] if fila_enc else []
            if "Máquina" not in enc or "Odómetro" not in enc:
                return []
            i_maq, i_odo = enc.index("Máquina"), enc.index("Odómetro")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) <= max(i_maq, i_odo):
                    continue
                maq = str(row[i_maq] or "").strip()
                if not maq or row[i_odo] in (None, ""):
                    continue                    # sin lectura no cuenta
                cuenta[maq] += 1
                fecha = str(row[0] or "")
                if fecha > ultima.get(maq, ""):
                    ultima[maq] = fecha
        finally:
            wb.close()
    except Exception as e:
        logger.warning("opciones_capataz: no pude contar las lecturas: %r", e)
        return []
    return sorted(cuenta, key=lambda m: (cuenta[m], ultima.get(m, "")),
                  reverse=True)[:tope]


def labores_frecuentes(excel_path: str | None = None,
                       tope: int = MAX_LABORES) -> list:
    """Las labores mas usadas de la bitacora, de mas a menos."""
    from collections import Counter

    from openpyxl import load_workbook

    from config import EXCEL_PATH
    ruta = excel_path or EXCEL_PATH
    cuenta, grafia = Counter(), {}
    try:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            if BITACORA_SHEET not in wb.sheetnames:
                return []
            for row in wb[BITACORA_SHEET].iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 4 or not row[3]:
                    continue
                act = str(row[3]).strip()
                clave = act.lower()
                if clave in _NO_ES_LABOR:
                    continue
                cuenta[clave] += 1
                grafia.setdefault(clave, act)   # la primera grafía que se vio
        finally:
            wb.close()
    except Exception as e:
        logger.warning("opciones_capataz: no pude leer las labores: %r", e)
        return []
    return [grafia[c] for c, _ in cuenta.most_common(tope)]
