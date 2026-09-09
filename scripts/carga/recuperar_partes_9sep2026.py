# -*- coding: utf-8 -*-
"""Reingresa los cinco partes que un horometro a medias se comio el 9-sep-2026.

QUE PASO
Probando el flujo guiado el 8-sep a las 16:23, Juan dejo un horometro sin
terminar: `horo_state='esperando_termino'` quedo en el pickle. Al dia siguiente
mando CINCO partes de asistencia --3, 4, 7, 8 y 9 de septiembre-- y se los comio
ese flujo, uno por uno, sin una sola linea en el log.

El guardia que caduca los flujos a los 30 min existe desde el 2-sep, pero cuando
el flujo guiado de horometro se agrego el 7-sep quedo POR ENCIMA de el en
handlers/chat.py: era el unico flujo inmune a caducar. La prueba esta en el
pickle, donde `flujo_ts` ni siquiera se creo. Arreglado moviendo revisar_flujos
antes del bloque del horometro.

Los mensajes NO se perdieron: estan crudos en files/telegram/2026-09.jsonl.

QUE HACE
El mismo camino que habria seguido el bot: parsear_asistencia_multi() y una
fila por labor, en orden cronologico. La fecha sale del propio texto ("Lunes 7
de septiembre 2026"), no de hoy: se esta reconstruyendo historia.

USO
    python scripts/carga/recuperar_partes_9sep2026.py --simular
    python scripts/carga/recuperar_partes_9sep2026.py
"""
import os
import shutil
import sys
from datetime import datetime

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)

# Con un prefijo para comprobar que cada mensaje es el que creemos ANTES de
# escribir nada.
PERDIDOS = [
    (3239, "Asistencia 3 de septiembre 2026"),
    (3242, "Asistencia 4 de septiembre 2026"),
    (3245, "Lunes 7 de septiembre 2026"),
    (3248, "Martes 8 de septiembre 2026"),
    (3251, "Miércoles 9 de septiembre 2026"),
]
QUIEN = "Juan Parada"


def _mensajes():
    from modules.telegram_backup import leer_mes
    por_id = {f.get("message_id"): f for f in leer_mes("2026-09") if f.get("text")}
    salida = []
    for mid, prefijo in PERDIDOS:
        fila = por_id.get(mid)
        if fila is None:
            raise SystemExit("No encontré el mensaje %s en el respaldo crudo." % mid)
        if not fila["text"].startswith(prefijo):
            raise SystemExit("El mensaje %s no es el esperado: %r"
                             % (mid, fila["text"][:60]))
        salida.append(fila)
    return salida


def _fecha_del_texto(texto):
    from modules.bitacora_asistencia import fecha_de_linea
    for linea in texto.splitlines()[:2]:
        f = fecha_de_linea(linea)
        if f:
            return f
    return None


def _filas_de(fila):
    """Traduce un mensaje crudo a filas de bitacora. (filas, motivo si no va)."""
    from modules.bitacora_asistencia import (cultivo_de, parsear_asistencia_multi,
                                             tipo_de)
    texto = fila["text"]
    fecha_txt = _fecha_del_texto(texto)
    lineas = [l for l in texto.splitlines()[1:] if l.strip()]

    dias = parsear_asistencia_multi(texto)
    if not dias:
        return [], "parsear_asistencia_multi no leyó ninguna línea"

    filas, personas = [], 0
    for d in dias:
        f = d["fecha"] or fecha_txt
        for g in d["grupos"]:
            personas += len(g["trabajadores"])
            filas.append({
                "fecha": f.strftime("%Y-%m-%d") if f else "",
                "tipo": tipo_de(g["actividad"]), "actividad": g["actividad"],
                "cultivo": cultivo_de(g["actividad"]), "sector": "",
                "jornadas_hombre": g["jornadas_hombre"],
                "trabajadores": g["trabajadores"], "insumo": "",
                "cantidad": None, "unidad": "", "maquina": "",
                "odometro": None, "superficie_ha": None,
                "texto_original": texto,
            })
    if personas < len(lineas):
        return [], ("se leyeron %d de %d líneas" % (personas, len(lineas)))
    return filas, None


def _ya_estan(filas):
    """Fechas que YA tienen filas de asistencia en la bitacora.

    La comprobacion va contra la fecha con la que se va a ESCRIBIR, no contra
    la que declara el mensaje. Por saltarse esto el 7-sep se duplicaron dos
    filas: el texto decia "2028" y la comparacion daba cero.
    """
    from openpyxl import load_workbook

    from config import EXCEL_PATH
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb["Bitácora"]
    enc = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i = {n: k for k, n in enumerate(enc)}
    existentes = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[i["Trabajadores"]]:
            existentes.add(str(r[0])[:10])
    wb.close()
    return {f["fecha"] for f in filas} & existentes


def main(simular):
    filas, pendientes = [], []
    for fila in _mensajes():
        f, motivo = _filas_de(fila)
        if motivo:
            pendientes.append((fila, motivo))
        filas.extend(f)
    filas.sort(key=lambda f: f["fecha"])

    print("%d filas a reingresar, de %d partes\n" % (len(filas), len(PERDIDOS)))
    for f in filas:
        print("  %s  %-11s %-32s %4s JH  %s"
              % (f["fecha"], f["tipo"], f["actividad"][:32],
                 f["jornadas_hombre"] or 0, ", ".join(f["trabajadores"])[:46]))
    print("\n  total jornadas-hombre: %g"
          % sum(f["jornadas_hombre"] or 0 for f in filas))

    if pendientes:
        print("\nNO se reingresan (siguen enteros en el respaldo crudo):")
        for fila, motivo in pendientes:
            print("  mid=%s  %-40s %s" % (fila["message_id"],
                                          fila["text"].splitlines()[0][:40], motivo))

    chocan = _ya_estan(filas)
    if chocan:
        print("\n⛔ ABORTA: estas fechas YA tienen asistencia anotada: %s"
              % ", ".join(sorted(chocan)))
        return 2

    if simular:
        print("\n--simular: no se escribió nada.")
        return 0

    from config import EXCEL_PATH
    resp = shutil.copy2(EXCEL_PATH, EXCEL_PATH.replace(
        ".xlsx", "_bak_%s.xlsx" % datetime.now().strftime("%Y%m%d_%H%M%S")))
    print("\nRespaldo: %s" % os.path.basename(resp))

    from bitacora_manager import registrar_bitacora_estructurada
    ok = 0
    for f in filas:
        registrar_bitacora_estructurada(f, QUIEN)
        ok += 1
    print("%d filas escritas." % ok)
    return 0


if __name__ == "__main__":
    sys.exit(main("--simular" in sys.argv))
