# -*- coding: utf-8 -*-
"""Corrige los saldos de vacaciones al 15-sep-2026.

Revisión del 15-sep-2026 (memoria: project_vacaciones_saldos_sep2026):
  · Los saldos de Personal eran EXACTAMENTE el cálculo de
    cargar_vacaciones_historico.py del 4-ago, con agosto incluido. Nadie los
    tocó después: el job mensual nunca acumuló con éxito.
  · Faltaba septiembre: +1,25 a cada uno. El job del 1-sep no corrió porque
    el PC estaba suspendido.
  · El Excel de Juan (VACAC 26, modificado el 11-sep) tiene dos vacaciones que
    el Master no: Patricio del 10 al 21-ago (10 días) y Javier del 20 al 21-ago
    (2 días). El dueño confirmó las de Javier el 15-sep: el resumen de Juan
    dice 9, pero el detalle (11) es el correcto.

⚠️ NO usa cargar_vacaciones_historico.py --aplicar: ese script borra la hoja
Vacaciones entera y recarga solo lo del Excel de Juan, y se perderían las filas
de Felix De Vicente, que no vienen de ahí.

Este solo AGREGA las filas que faltan y actualiza, para los 6 de Personal,
Días Pendientes, Días Tomados Total, Última Vacación y Notas.

Uso (desde la carpeta Robot):
  python scripts/carga/corregir_vacaciones_sep2026.py            simulación
  python scripts/carga/corregir_vacaciones_sep2026.py --aplicar  con el bot DETENIDO
"""
import base64
import hashlib
import os
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date, datetime

ROBOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROBOT)
sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook  # noqa: E402

import config  # noqa: E402

APLICAR = "--aplicar" in sys.argv
MES_ACUMULADO = (2026, 9)       # suma la acumulación hasta septiembre 2026, inclusive
MES_DEL_4AGO = (2026, 8)        # el cálculo del 4-ago llegaba hasta agosto
DIAS_MES = 15.0 / 12.0
NOTA_FILA = "Periodo 2026 · del Excel de Juan, agregado el 15-sep-2026"

ORIGEN = os.path.join(os.path.dirname(config.DROPBOX_BASE), "CAMARICO 2023",
                      "ASISTENCIA TEMP 2023-2024 fda JUAN PARADA "
                      "(Copia en conflicto de juan parada 2025-03-06).xlsx")

# Las dos vacaciones revisadas y confirmadas. Si el Excel de Juan trae otras
# que el Master no tiene, el script se detiene: no las revisó nadie.
CONFIRMADAS = {
    ("Luis Patricio Mora Amigo", date(2026, 8, 10), date(2026, 8, 21), 10.0),
    ("Javier Gonzalez", date(2026, 8, 20), date(2026, 8, 21), 2.0),
}

CANON = {
    "MORA PATRICIO": "Luis Patricio Mora Amigo",
    "PARADA JUAN": "Juan Parada Castillo",
    "AMIGO RAMIRO": "Luis Ramiro Amigo Soto",
    "RAMIRO AMIGO": "Luis Ramiro Amigo Soto",
    "AMIGO FELICITO": "Felicito Amigo Soto",
    "FELICITO AMIGO": "Felicito Amigo Soto",
    "MORA AGUSTIN": "Agustin Segundo Mora Hernandez",
    "AGUSTIN MORA": "Agustin Segundo Mora Hernandez",
    "GONZALES JAVIER": "Javier Gonzalez",
    "GONZALEZ JAVIER": "Javier Gonzalez",
}


def norm(n):
    return " ".join(str(n or "").upper().split())


def dia(v):
    return v.date() if isinstance(v, datetime) else v


def corregir_anio(f, anio):
    """Igual que la carga del 4-ago: el año del bloque manda."""
    if not hasattr(f, "year"):
        return None
    f = dia(f)
    if f.year == anio:
        return f
    try:
        return f.replace(year=anio)
    except ValueError:
        return f


def leer_excel_juan():
    wb = load_workbook(ORIGEN, read_only=True, data_only=True)
    filas = list(wb["VACAC 26"].iter_rows(values_only=True))
    wb.close()
    registros, anio = [], None
    for row in filas:
        if not row:
            continue
        m = re.match(r"AÑO\s*(\d{2})", str(row[0] or "").strip().upper())
        if m:
            anio = 2000 + int(m.group(1))
            continue
        nombre = CANON.get(norm(row[2])) if len(row) > 2 else None
        if not nombre or anio is None:
            continue
        try:
            dias = float(row[5] or 0)
        except (TypeError, ValueError):
            continue
        if dias > 0:
            registros.append((nombre, corregir_anio(row[3], anio),
                              corregir_anio(row[4], anio), dias))
    return registros


def saldo(base, tomados, mes):
    s, f0 = base
    meses = (mes[0] - f0.year) * 12 + (mes[1] - f0.month)
    return round(s + meses * DIAS_MES - tomados, 2)


def huella(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def procesos_del_bot():
    ps = ("(Get-CimInstance Win32_Process -Filter \"Name='python3.11.exe'\" | "
          "Where-Object { $_.CommandLine -like '*main.py*' } | Measure-Object).Count")
    r = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-EncodedCommand",
         base64.b64encode(ps.encode("utf-16-le")).decode()],
        capture_output=True, text=True, timeout=60)
    salida = (r.stdout or "").strip()
    if r.returncode != 0 or not salida.isdigit():
        sys.exit("No pude comprobar si el bot está corriendo: no escribí nada.")
    return int(salida)


def respaldar():
    """Copia verificada a Dropbox\\Backups y a la cola de Drive, SIN rotar.

    No usa backup_master: su rotación borra por orden alfabético y se llevaría
    un respaldo real (ver project_revision_raspberry_pi).
    """
    from modules.drive.cola import Cola
    carpeta = os.path.join(config.DROPBOX_BACKUP_PATH, "Master")
    snaps = os.path.join(carpeta, "snapshots")
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M")
    snap = os.path.join(snaps, ts + ".xlsx")
    parcial = os.path.join(snaps, ts + ".parcial.xlsx")
    if os.path.exists(snap):
        sys.exit(f"Ya existe {snap}: espera un minuto y vuelve a correr.")
    antes = huella(config.EXCEL_PATH)
    shutil.copy2(config.EXCEL_PATH, parcial)
    if huella(parcial) != antes or huella(config.EXCEL_PATH) != antes:
        os.remove(parcial)
        sys.exit("El Master cambió durante el respaldo: no escribí nada.")
    wb = load_workbook(parcial, read_only=True)
    sano = {"Personal", "Vacaciones"} <= set(wb.sheetnames)
    wb.close()
    if not sano:
        sys.exit(f"El respaldo no se ve sano (quedó en {parcial}): no escribí nada.")
    os.replace(parcial, snap)
    actual = os.path.join(carpeta, "current.parcial.xlsx")
    shutil.copy2(snap, actual)
    os.replace(actual, os.path.join(carpeta, "current.xlsx"))
    Cola(config.DRIVE_COLA_PATH, config.DRIVE_MAX_INTENTOS).encolar(
        snap, "Respaldos/Master", os.path.basename(snap))
    return snap


def leer_master(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    vac = [(str(r[0]).strip(), dia(r[1]), dia(r[2]), float(r[3] or 0))
           for r in wb["Vacaciones"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    base = {str(r[0]).strip(): (float(r[3] or 0), dia(r[4]))
            for r in wb["Vacaciones Pendientes"].iter_rows(min_row=2, values_only=True)
            if r and r[0]}
    personal = {}
    for r in wb["Personal"].iter_rows(min_row=2, values_only=True):
        if r and r[0]:
            personal[str(r[0]).strip()] = {"pend": float(r[4] or 0),
                                           "tom": float(r[5] or 0),
                                           "ult": dia(r[6])}
    hojas = len(wb.sheetnames)
    wb.close()
    return vac, base, personal, hojas


def main():
    hoy = date.today()
    if (hoy.year, hoy.month) != MES_ACUMULADO:
        sys.exit(f"Este ajuste suma la acumulación hasta {MES_ACUMULADO[1]:02d}-"
                 f"{MES_ACUMULADO[0]} y hoy es {hoy}: hay que revisar de nuevo.")

    juan = leer_excel_juan()
    vac, base, personal, hojas = leer_master(config.EXCEL_PATH)

    # ── ¿Qué trae el Excel de Juan que el Master no tiene? ──
    nuevas = set((Counter(juan) - Counter(vac)).keys())
    sin_revisar = nuevas - CONFIRMADAS
    if sin_revisar:
        print("⛔ El Excel de Juan trae vacaciones que el Master no tiene y nadie revisó:")
        for n, a, b, d in sorted(sin_revisar, key=lambda x: (x[0], str(x[1]))):
            print(f"   {n}  {a} a {b}  {d:g} días")
        sys.exit("No escribí nada.")
    a_agregar = sorted(nuevas & CONFIRMADAS, key=lambda x: (x[1], x[0]))

    tom_ahora, tom_despues, ultima = defaultdict(float), defaultdict(float), {}
    for n, _a, _b, d in vac:
        tom_ahora[n] += d
    for n, _a, b, d in vac + a_agregar:
        tom_despues[n] += d
        if isinstance(b, date) and (n not in ultima or b > ultima[n]):
            ultima[n] = b

    # ── Plan, con el "¿ya está?" contra lo que se va a ESCRIBIR ──
    plan, problemas = [], []
    for n, p in personal.items():
        if n not in base:
            problemas.append(f"{n}: no está en la hoja Vacaciones Pendientes")
            continue
        objetivo = saldo(base[n], tom_despues[n], MES_ACUMULADO)
        del_4ago = saldo(base[n], tom_ahora[n], MES_DEL_4AGO)
        if abs(p["pend"] - objetivo) < 0.005 and abs(p["tom"] - tom_despues[n]) < 0.005:
            estado = "ya estaba"
        elif abs(p["pend"] - del_4ago) < 0.005 and abs(p["tom"] - tom_ahora[n]) < 0.005:
            estado = "corregir"
        else:
            problemas.append(f"{n}: tiene {p['pend']:.2f} pendientes y {p['tom']:.0f} tomados; "
                             f"no es el cálculo del 4-ago ({del_4ago:.2f}) ni el corregido "
                             f"({objetivo:.2f}). Alguien lo cambió: revisar a mano.")
            continue
        plan.append((n, p, objetivo, tom_despues[n], ultima.get(n, p["ult"]), estado))
    if problemas:
        print("⛔ " + "\n⛔ ".join(problemas))
        sys.exit("No escribí nada.")

    print(f"{'trabajador':32} {'días pendientes':>17}  {'tomados':>9}  "
          f"{'última vacación':>25}  estado")
    for n, p, obj, tom, ult, estado in plan:
        print(f"{n:32} {p['pend']:>7.2f} → {obj:>6.2f}  {p['tom']:>3.0f} → {tom:>3.0f}  "
              f"{str(p['ult']):>11} → {str(ult):>10}  {estado}")
    if a_agregar:
        print("\nFilas nuevas en la hoja Vacaciones:")
        for n, a, b, d in a_agregar:
            print(f"  + {n}  {a} a {b}  {d:g} días  «{NOTA_FILA}»")
    else:
        print("\nFilas nuevas en la hoja Vacaciones: ninguna")

    por_corregir = [x for x in plan if x[5] == "corregir"]
    if not por_corregir and not a_agregar:
        print("\n✅ Ya estaba aplicado: no hay nada que escribir.")
        return
    if not APLICAR:
        print("\n(simulación: no se escribió nada. Para aplicar, detener el bot y "
              "correr con --aplicar)")
        return

    # ── Aplicar ──
    n_bot = procesos_del_bot()
    if n_bot:
        sys.exit(f"⛔ El bot está corriendo ({n_bot} proceso): detenlo antes de aplicar. "
                 f"No escribí nada.")
    bloqueo = os.path.join(os.path.dirname(config.EXCEL_PATH),
                           "~$" + os.path.basename(config.EXCEL_PATH))
    if os.path.exists(bloqueo):
        sys.exit("⛔ El Master está abierto en Excel: ciérralo antes de aplicar. "
                 "No escribí nada.")

    snap = respaldar()
    print(f"\nRespaldo previo: {os.path.basename(snap)} (encolado para Drive)")

    wb = load_workbook(config.EXCEL_PATH)          # sin data_only: conserva las fórmulas
    ws_vac, ws_per = wb["Vacaciones"], wb["Personal"]
    for n, a, b, d in a_agregar:
        ws_vac.append([n, a, b, int(d) if float(d).is_integer() else d, "Aprobado", NOTA_FILA])
    filas = {str(ws_per.cell(r, 1).value or "").strip(): r for r in range(2, ws_per.max_row + 1)}
    for n, _p, obj, tom, ult, _estado in por_corregir:
        r = filas[n]
        ws_per.cell(r, 5).value = obj
        ws_per.cell(r, 6).value = int(tom) if float(tom).is_integer() else tom
        ws_per.cell(r, 7).value = ult
        ws_per.cell(r, 8).value = (f"Histórico 2023-2026 del Excel de Juan (VACAC 26) · "
                                   f"acumulado hasta 09-2026 · {tom:.0f} días tomados · "
                                   f"corregido el 15-sep-2026")

    # Guardado atómico: a un temporal que se verifica ANTES de tocar el Master
    temporal = os.path.join(os.path.dirname(config.EXCEL_PATH),
                            f"MASTER_corrigiendo_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb.save(temporal)
    wb.close()

    def comprobar(path):
        v2, _b, per2, hojas2 = leer_master(path)
        errores = []
        if hojas2 != hojas:
            errores.append(f"{hojas2} hojas en vez de {hojas}")
        if len(v2) != len(vac) + len(a_agregar):
            errores.append(f"Vacaciones tiene {len(v2)} filas en vez de "
                           f"{len(vac) + len(a_agregar)}")
        for n, _p, obj, tom, _ult, _estado in plan:
            q = per2.get(n)
            if not q or abs(q["pend"] - obj) >= 0.005 or abs(q["tom"] - tom) >= 0.005:
                errores.append(f"{n}: quedó {q}")
        return errores

    errores = comprobar(temporal)
    if errores:
        sys.exit(f"⛔ La copia corregida no quedó bien ({'; '.join(errores)}). "
                 f"El Master NO se tocó; la copia quedó en {temporal}.")
    try:
        os.replace(temporal, config.EXCEL_PATH)
    except PermissionError:
        sys.exit(f"⛔ No pude reemplazar el Master (¿abierto?). NO se tocó; "
                 f"la corrección quedó en {temporal}.")
    errores = comprobar(config.EXCEL_PATH)
    if errores:
        sys.exit(f"⛔ El Master quedó distinto a lo planeado: {'; '.join(errores)}. "
                 f"Restaurar desde {snap}.")
    print(f"✅ Aplicado y verificado en el Master: {len(a_agregar)} filas nuevas en "
          f"Vacaciones, {len(por_corregir)} trabajadores actualizados.")


if __name__ == "__main__":
    main()
