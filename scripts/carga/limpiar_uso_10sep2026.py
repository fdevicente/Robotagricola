# -*- coding: utf-8 -*-
"""Limpia lo que dejo el /uso mal cerrado del 10-sep-2026.

QUE PASO
Juan abrio /uso y dijo "Katana", 5 kilos, Nogales. Como `_find_producto` exigia
coincidencia EXACTA, "Katana" no calzo con "KATANA 1 KG - herbicida" y se creo
una fila NUEVA con stock -5. Despues escribio "/ uso" CON ESPACIO --que Telegram
no marca como comando-- asi que "Ripper full" y "60 litros en nogales" cayeron
como texto libre en la bitacora en vez de entrar al flujo.

Las dos causas ya estan arregladas (buscar_producto y comando_con_espacio).
Esto limpia lo que quedo escrito.

LO QUE HACE, QUE ES SOLO LO INEQUIVOCO
  - Borra la fila de inventario con el producto EN BLANCO y su fila de
    Aplicaciones con cantidad 0.
  - Funde "Katana" (-5 L, fantasma) en "KATANA 1 KG - herbicida": descuenta ahi
    los 5 que Juan uso y borra la fantasma. Juan dijo "5 KILOS" y ese producto
    esta en Kg, asi que tambien se corrige la unidad en Aplicaciones.

LO QUE NO TOCA, A PROPOSITO
  - "Agrocupper" en -2,1 L (6-ago). El producto real es "Agrocupper SP", en Kg:
    cambia el nombre Y la unidad. Eso lo decide el dueño.
  - Los 60 L de Ripper Full. Hay TRES filas candidatas (20 L, 80 L y 60 L) y
    elegir una al azar es peor que dejarlo.
  - Las dos filas de bitacora tipificadas como MAQUINARIA.

USO
    python scripts/carga/limpiar_uso_10sep2026.py --simular
    python scripts/carga/limpiar_uso_10sep2026.py
"""
import os
import shutil
import sys
from datetime import datetime

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)

FANTASMA = "Katana"                  # la fila que se creo sola
REAL = "KATANA 1 KG - herbicida"     # el producto que existia
USADO = 5


def _fila_de(ws, nombre):
    for n in range(2, ws.max_row + 1):
        if str(ws.cell(n, 1).value or "").strip() == nombre:
            return n
    return None


def main(simular):
    from openpyxl import load_workbook

    from config import EXCEL_PATH
    from excel_manager import _save_wb

    wb = load_workbook(EXCEL_PATH)
    wi, wa = wb["Inventario"], wb["Aplicaciones"]

    f_fant = _fila_de(wi, FANTASMA)
    f_real = _fila_de(wi, REAL)
    vacias_inv = [n for n in range(2, wi.max_row + 1)
                  if not str(wi.cell(n, 1).value or "").strip()
                  and any(wi.cell(n, c).value not in (None, "") for c in range(2, 8))]
    vacias_app = [n for n in range(2, wa.max_row + 1)
                  if wa.cell(n, 1).value
                  and not str(wa.cell(n, 2).value or "").strip()]

    # La comprobacion va contra lo que se va a ESCRIBIR, no contra lo que uno
    # cree recordar: si el fantasma ya no esta, no hay nada que fundir.
    if f_fant is None:
        print("La fila fantasma %r ya no está. Nada que fundir." % FANTASMA)
    elif f_real is None:
        print("⛔ ABORTA: no encuentro %r en el inventario." % REAL)
        return 2
    else:
        neg = wi.cell(f_fant, 4).value
        stock_real = float(wi.cell(f_real, 4).value or 0)
        print("FUNDIR")
        print("  fila %-4d %-36s stock %s   <- se borra" % (f_fant, FANTASMA, neg))
        print("  fila %-4d %-36s stock %g -> %g"
              % (f_real, REAL, stock_real, stock_real - USADO))
        if float(neg or 0) != -USADO:
            print("  ⛔ ABORTA: esperaba %g en la fantasma y encontré %s"
                  % (-USADO, neg))
            return 2

    print("\nBORRAR FILAS EN BLANCO")
    for n in vacias_inv:
        print("  Inventario fila %d" % n)
    for n in vacias_app:
        print("  Aplicaciones fila %d" % n)
    if not vacias_inv and not vacias_app:
        print("  ninguna")

    print("\nCORREGIR LA UNIDAD EN APLICACIONES")
    a_corregir = [n for n in range(2, wa.max_row + 1)
                  if str(wa.cell(n, 2).value or "").strip() == FANTASMA]
    for n in a_corregir:
        print("  fila %-4d %s: %s -> Kg, producto -> %s"
              % (n, wa.cell(n, 1).value, wa.cell(n, 4).value, REAL))
    if not a_corregir:
        print("  ninguna")

    if simular:
        print("\n--simular: no se escribió nada.")
        return 0

    resp = shutil.copy2(EXCEL_PATH, EXCEL_PATH.replace(
        ".xlsx", "_bak_%s.xlsx" % datetime.now().strftime("%Y%m%d_%H%M%S")))
    print("\nRespaldo: %s" % os.path.basename(resp))

    for n in a_corregir:
        wa.cell(n, 2).value = REAL
        wa.cell(n, 4).value = "Kg"
    if f_fant is not None and f_real is not None:
        wi.cell(f_real, 4).value = float(wi.cell(f_real, 4).value or 0) - USADO
        wi.cell(f_real, 7).value = wi.cell(f_fant, 7).value
    # De abajo hacia arriba, para que no se corran los indices.
    for n in sorted(vacias_app, reverse=True):
        wa.delete_rows(n, 1)
    for n in sorted(([f_fant] if f_fant else []) + vacias_inv, reverse=True):
        wi.delete_rows(n, 1)

    _save_wb(wb)
    print("Listo.")
    return 0


if __name__ == "__main__":
    sys.exit(main("--simular" in sys.argv))
