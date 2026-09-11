# -*- coding: utf-8 -*-
"""Unifica los productos repetidos del inventario y cierra el /uso del 10-sep.

EL DUEÑO, 11-sep-2026: "deberian tomar por el nombre, y unificarlos en uno solo.
si no calza preguntar". Y sobre el stock: "segun lo que se sume en factura no
deberia importar si estan unificados" — o sea que una vez unificados, da lo
mismo de cual fila se descuenta.

QUE UNIFICA
Dos grupos, medidos contra el Master:
  - Ripper Full: 3 filas (20 + 80 + 60 = 160 L), misma unidad y categoria.
  - Agrocupper:  2 filas. "Agrocupper SP" (Fungicida, 4,9 Kg) y el fantasma
    "Agrocupper" (Otro, -2,1 L) que nacio del buscador exacto.

⚠️ LOS CUATRO DEFENDER NO SE TOCAN. "Defender Zn", "Defender Calcio",
"Defender K (Potasio)" y "Defender Boro" comparten la primera palabra y son
productos DISTINTOS. Por eso la regla es PREFIJO DE PALABRAS y no primera
palabra: ver agrupar_duplicados().

QUE MAS HACE
  - Descuenta los 60 L de Ripper Full que Juan uso el 10-sep y que nunca se
    anotaron: su mensaje cayo como texto libre porque escribio "/ uso" con
    espacio. Quedan 100 L.
  - Borra las dos filas de bitacora que ese mismo error dejo tipificadas como
    MAQUINARIA ("Ripper full" con maquina TRACTOR y "Trabajo con ripper" con
    maquina RIPPER): son los mensajes mal ruteados, no trabajo real.
  - Completa la fila de APLICACION del 10-sep, que anotaba 60 L sin decir de que.

COMO ELIGE EL NOMBRE QUE SOBREVIVE
El mas corto DE ENTRE los que tienen una categoria de verdad. El fantasma
siempre queda con categoria "Otro", asi que su nombre --que es la abreviatura
que escribio Juan-- no puede ganarle al comercial. Para Ripper Full los tres son
legitimos y gana "Ripper Full", que es justo como lo escribe el.

USO
    python scripts/carga/unificar_inventario_11sep2026.py --simular
    python scripts/carga/unificar_inventario_11sep2026.py
"""
import os
import shutil
import sys
from datetime import datetime

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)

USO_RIPPER = 60
FECHA_USO = "2026-09-10"
CULTIVO_USO = "NOGALES"


def _leer_inventario(ws):
    filas = []
    for n in range(2, ws.max_row + 1):
        if ws.cell(n, 1).value:
            filas.append({
                "fila": n, "nombre": str(ws.cell(n, 1).value).strip(),
                "categoria": ws.cell(n, 2).value, "unidad": ws.cell(n, 3).value,
                "stock": float(ws.cell(n, 4).value or 0),
                "minimo": ws.cell(n, 5).value,
                "entrada": ws.cell(n, 6).value, "uso": ws.cell(n, 7).value,
            })
    return filas


def _superviviente(grupo):
    """La fila que se queda: el nombre mas corto con categoria de verdad."""
    reales = [g for g in grupo
              if str(g["categoria"] or "").strip().lower() not in ("", "otro")]
    return min(reales or grupo, key=lambda g: len(g["nombre"]))


def main(simular):
    from openpyxl import load_workbook

    from config import EXCEL_PATH
    from excel_manager import _save_wb
    from inventario_manager import agrupar_duplicados

    wb = load_workbook(EXCEL_PATH)
    wi, wa, wb_bit = wb["Inventario"], wb["Aplicaciones"], wb["Bitácora"]
    filas = _leer_inventario(wi)
    grupos = agrupar_duplicados(filas)

    print("PRODUCTOS A UNIFICAR (%d grupos)\n" % len(grupos))
    plan = []
    for g in grupos:
        queda = _superviviente(g)
        suma = sum(x["stock"] for x in g)
        unidades = {str(x["unidad"]).strip() for x in g}
        print("  → %s  ·  %g %s" % (queda["nombre"], suma, queda["unidad"]))
        for x in g:
            marca = "SE QUEDA" if x is queda else "se funde"
            print("      %-8s fila %-4d %-42s %-4s %g"
                  % (marca, x["fila"], x["nombre"][:42], x["unidad"], x["stock"]))
        if len(unidades) > 1:
            print("      ⚠️ unidades distintas %s — se usa la del que se queda (%s)"
                  % (sorted(unidades), queda["unidad"]))
        plan.append((queda, [x for x in g if x is not queda], suma))
    print()

    # ── El uso de Ripper Full que nunca se anotó ──
    ripper = next((q for q, _, _ in plan
                   if "ripper" in q["nombre"].lower()), None)
    if ripper:
        print("DESCONTAR EL USO DEL 10-SEP")
        total = next(s for q, _, s in plan if q is ripper)
        print("  %s: %g → %g %s  (%g usados en %s el %s)\n"
              % (ripper["nombre"], total, total - USO_RIPPER, ripper["unidad"],
                 USO_RIPPER, CULTIVO_USO, FECHA_USO))

    # ── Las filas de bitácora mal ruteadas ──
    borrar_bit, completar_bit = [], None
    enc = [c.value for c in next(wb_bit.iter_rows(min_row=1, max_row=1))]
    i = {n: k for k, n in enumerate(enc)}
    for n in range(2, wb_bit.max_row + 1):
        if str(wb_bit.cell(n, i["Fecha"] + 1).value)[:10] != FECHA_USO:
            continue
        tipo = str(wb_bit.cell(n, i["Tipo"] + 1).value or "")
        act = str(wb_bit.cell(n, i["Actividad"] + 1).value or "")
        if tipo == "MAQUINARIA" and "ripper" in act.lower():
            borrar_bit.append((n, act))
        elif tipo == "APLICACION" and not wb_bit.cell(n, i["Insumo"] + 1).value:
            completar_bit = n
    print("BITÁCORA DEL 10-SEP")
    for n, act in borrar_bit:
        print("  borrar     fila %-4d MAQUINARIA  %s" % (n, act))
    if completar_bit:
        print("  completar  fila %-4d APLICACION  insumo → %s"
              % (completar_bit, ripper["nombre"] if ripper else "?"))
    if not borrar_bit and not completar_bit:
        print("  nada que hacer")

    if simular:
        print("\n--simular: no se escribió nada.")
        return 0

    resp = shutil.copy2(EXCEL_PATH, EXCEL_PATH.replace(
        ".xlsx", "_bak_%s.xlsx" % datetime.now().strftime("%Y%m%d_%H%M%S")))
    print("\nRespaldo: %s" % os.path.basename(resp))

    # Inventario: el que se queda toma la suma; los otros se borran de abajo
    # hacia arriba para que no se corran los índices.
    a_borrar = []
    for queda, funden, suma in plan:
        total = suma - (USO_RIPPER if queda is ripper else 0)
        wi.cell(queda["fila"], 4).value = total
        if queda is ripper:
            wi.cell(queda["fila"], 7).value = FECHA_USO
        a_borrar += [x["fila"] for x in funden]
    for n in sorted(a_borrar, reverse=True):
        wi.delete_rows(n, 1)

    if ripper:
        wa.append([FECHA_USO, ripper["nombre"], USO_RIPPER, ripper["unidad"],
                   CULTIVO_USO, "", "Juan Parada",
                   "Uso reingresado: su mensaje cayó como texto libre "
                   "porque escribió '/ uso' con espacio"])
        if completar_bit:
            wb_bit.cell(completar_bit, i["Insumo"] + 1).value = ripper["nombre"]
    for n, _ in sorted(borrar_bit, reverse=True):
        wb_bit.delete_rows(n, 1)

    _save_wb(wb)
    print("Listo: %d filas de inventario fundidas, %d de bitácora borradas."
          % (len(a_borrar), len(borrar_bit)))
    return 0


if __name__ == "__main__":
    sys.exit(main("--simular" in sys.argv))
